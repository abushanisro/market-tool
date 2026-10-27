"""Downloadable document parsers: CSV, XLSX, (permitted) PDF tables, ISO open data."""
from __future__ import annotations

import csv
import io

from . import ParseContext, register
from .generic import rows_to_items


def _csv_rows(text: str) -> tuple[list[str], list[list[str]]]:
    rdr = csv.reader(io.StringIO(text))
    header = next(rdr, [])
    return [h.strip().lstrip("﻿") for h in header], [r for r in rdr if any(c.strip() for c in r)]


@register("csv_table")
def csv_table(response, ctx: ParseContext):
    header, rows = _csv_rows(response.text)
    yield from rows_to_items(header, rows, ctx, response.url)


@register("xlsx_table")
def xlsx_table(response, ctx: ParseContext):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.body), read_only=True, data_only=True)
    ws = wb[ctx.options["sheet"]] if ctx.options.get("sheet") else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr_row = int(ctx.options.get("header_row", 0))
    for _ in range(hdr_row):
        next(it, None)
    header = [str(c).strip() if c is not None else "" for c in next(it, [])]
    rows = [list(r) for r in it if r and any(c not in (None, "") for c in r)]
    yield from rows_to_items(header, rows, ctx, response.url)


@register("pdf_tables")
def pdf_tables(response, ctx: ParseContext):
    """Only for PDFs whose licence permits extraction (enforced by source config review)."""
    import pdfplumber

    with pdfplumber.open(io.BytesIO(response.body)) as pdf:
        pages = ctx.options.get("pages") or range(len(pdf.pages))
        for p in pages:
            for table in pdf.pages[p].extract_tables():
                if not table or len(table) < 2:
                    continue
                header = [str(c or "").strip() for c in table[0]]
                if ctx.options.get("columns", {}).get("value") not in header:
                    continue
                yield from rows_to_items(header, [list(r) for r in table[1:]], ctx, response.url)


@register("iso_deliverables")
def iso_deliverables(response, ctx: ParseContext):
    """ISO Open Data deliverables metadata (ODC-By). Column names are looked up defensively."""
    header, rows = _csv_rows(response.text)
    low = {h.lower(): i for i, h in enumerate(header)}

    def col(*names):
        for n in names:
            if n.lower() in low:
                return low[n.lower()]
        return None

    i_ref = col("reference", "deliverableReference", "deliverable_reference")
    i_title = col("title.en", "title_en", "title")
    i_ics = col("icsCode", "ics", "ics_code")
    i_status = col("currentStage", "status", "publicationStatus")
    i_url = col("url", "link")
    if i_ref is None:
        raise ValueError(f"ISO deliverables CSV: reference column not found in {header[:12]}")
    prefixes = tuple(ctx.options.get("ics_prefixes", []))
    for r in rows:
        ics = r[i_ics] if i_ics is not None and i_ics < len(r) else ""
        if prefixes and not any(code.strip().startswith(prefixes) for code in ics.replace("|", ",").split(",")):
            continue
        yield {"item_type": "standard", "code": r[i_ref].strip(), "body": "ISO",
               "title": r[i_title] if i_title is not None else None,
               "status": r[i_status] if i_status is not None else None,
               "url": r[i_url] if i_url is not None else None, "raw": {"ics": ics}}
