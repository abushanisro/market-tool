"""World Bank Commodity Price Data ("Pink Sheet") — monthly historical XLSX.

Layout of the 'Monthly Prices' sheet: a few title rows, one row of commodity names, one row of
units such as '($/mt)', then rows keyed '1960M01' … '2026M09'. The parser locates those rows
instead of relying on fixed offsets, because the workbook layout has shifted over the years.
"""
from __future__ import annotations

import io
import re

from . import Follow, ParseContext, obs, register

PERIOD = re.compile(r"^\d{4}M\d{2}$")
MATERIAL_ALIASES = {"Aluminum": "Aluminium (LME benchmark)", "Copper": "Copper (LME benchmark)",
                    "Lead": "Lead (LME benchmark)", "Tin": "Tin (LME benchmark)", "Nickel": "Nickel (LME benchmark)",
                    "Zinc": "Zinc (LME benchmark)", "Iron ore, cfr spot": "Iron ore (CFR China)",
                    "Rubber, TSR20": "Natural rubber TSR20", "Rubber, RSS3": "Natural rubber RSS3"}


@register("worldbank_pinksheet")
def worldbank_pinksheet(response, ctx: ParseContext):
    ctype = (response.headers.get("Content-Type") or b"").decode().lower()
    if not (response.url.lower().endswith(".xlsx") or "spreadsheet" in ctype or response.body[:2] == b"PK"):
        # landing page → discover the current XLSX link (hash in the URL changes monthly)
        for href in response.css("a::attr(href)").getall():
            if re.search(r"CMO-Historical-Data-Monthly\.xlsx$", href, re.I):
                yield Follow(response.urljoin(href), parser="worldbank_pinksheet")
                return
        raise ValueError("Pink Sheet XLSX link not found on landing page")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.body), read_only=True, data_only=True)
    ws = wb[ctx.options.get("sheet", "Monthly Prices")]
    rows = list(ws.iter_rows(values_only=True))
    first_data = next(i for i, r in enumerate(rows) if r and isinstance(r[0], str) and PERIOD.match(r[0].strip()))
    unit_row = rows[first_data - 1]
    name_row = rows[first_data - 2]
    wanted = set(ctx.options.get("commodities") or [])
    keep = int(ctx.options.get("keep_months", 0)) or None
    data_rows = rows[first_data:]
    if keep:
        data_rows = data_rows[-keep:]
    for col in range(1, len(name_row)):
        name = (name_row[col] or "").strip() if isinstance(name_row[col], str) else None
        if not name or (wanted and name not in wanted):
            continue
        unit = str(unit_row[col] or "").strip("() ")
        for r in data_rows:
            v = r[col] if col < len(r) else None
            if v in (None, "", "…", ".."):
                continue
            yield obs(kind="price", entity=name, metric="monthly_average_price", value=v, unit=unit,
                      currency="USD" if "$" in unit else None, country="WLD", period_date=r[0].strip(),
                      material_name=MATERIAL_ALIASES.get(name, name),
                      original={"sheet": ws.title, "column": col, "header": name, "unit": unit, "period": r[0]})
